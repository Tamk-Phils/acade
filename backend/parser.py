"""
Enhanced Document Parser for extracting text, structure, tables, and metadata
from DOCX and PDF files with pinpoint accuracy.
"""
import re
import os
import docx
from typing import Dict, List, Any, Optional
from pypdf import PdfReader
from backend.models import DocumentMetadata
from backend.academic_data import UBA_ESTABLISHMENTS, resolve_department_and_option

IGNORE_TITLE_PHRASES = [
    "THE UNIVERSITY OF BAMENDA",
    "UNIVERSITE DE BAMENDA",
    "COLLEGE OF TECHNOLOGY",
    "REPUBLIC OF CAMEROON",
    "REPUBLIQUE DU CAMEROUN",
    "FACULTY OF SCIENCE",
    "FACULTY OF ARTS",
    "FACULTY OF EDUCATION",
    "FACULTY OF ECONOMICS",
    "PEACE - WORK - FATHERLAND",
    "COVER PAGE",
    "TITLE PAGE"
]

def clean_title(candidate: str) -> Optional[str]:
    if not candidate:
        return None
    cleaned = candidate.strip().strip('“”"\' \n\t')
    cleaned_upper = cleaned.upper()
    for ignore in IGNORE_TITLE_PHRASES:
        if ignore in cleaned_upper:
            return None
    if len(cleaned) < 6:
        return None
    return cleaned

def extract_metadata_from_text_and_tables(text: str, tables_text: List[str] = None) -> DocumentMetadata:
    """Extract metadata using high-precision legal phrases, tables, and regex patterns."""
    meta = DocumentMetadata()

    # 1. High-Precision Title Extraction
    # Rule A: Look inside statutory declaration, certification, or acceptance quotes
    decl_title_match = re.search(r'(?:titled|entitled)\s*[“"«]([^”"»\n]+)[”"»]', text, re.IGNORECASE)
    if decl_title_match:
        t = clean_title(decl_title_match.group(1))
        if t:
            meta.title = t

    # Rule B: Check title table cell if available (common in UBa documents)
    if (not meta.title or meta.title == "TITLE OF THE WORK") and tables_text:
        for cell_txt in tables_text:
            cleaned = clean_title(cell_txt)
            if cleaned and len(cleaned) >= 15 and len(cleaned) <= 220:
                meta.title = cleaned
                break

    # Rule C: Fallback to searching before purpose clause
    if not meta.title or meta.title == "TITLE OF THE WORK":
        before_clause = re.search(r'([A-Z0-9\s:,\-]{15,220})\n+\s*(?:A Dissertation|An Internship Report|A Project|A Final Year Project|A Thesis|A Research Proposal)', text, re.IGNORECASE)
        if before_clause:
            lines = [l.strip() for l in before_clause.group(1).split('\n') if l.strip()]
            for l in reversed(lines):
                cleaned = clean_title(l)
                if cleaned:
                    meta.title = cleaned
                    break

    # 2. Author Name
    author_match = re.search(r'(?:BY|PRESENTED BY|AUTHOR|CANDIDATE|BY:)\s*\n*([A-Z][a-zA-Z\s\-]+)', text, re.IGNORECASE)
    if author_match:
        cand = author_match.group(1).strip().split('\n')[0].strip()
        if len(cand) > 3 and "SUPERVISOR" not in cand.upper() and "UNIVERSITY" not in cand.upper():
            meta.author = cand

    # Check declaration for "I, [Name], registration N"
    decl_author = re.search(r'I,\s*([A-Z\s\-]{4,50}),\s*(?:registration|matricule)', text, re.IGNORECASE)
    if decl_author:
        cand = decl_author.group(1).strip()
        if len(cand) > 3 and "UNIVERSITY" not in cand.upper():
            meta.author = cand

    # 3. Registration Number
    reg_match = re.search(r'(?:REGISTRATION NUMBER|REG NO|MATRICULE|REGISTRATION N[°o◦]?)\s*[:.]?\s*([A-Za-z0-9\-]+)', text, re.IGNORECASE)
    if reg_match:
        meta.reg_number = reg_match.group(1).strip()

    # 4. Supervisor & Rank
    sup_match = re.search(r'(?:SUPERVISOR\(S\)?|SUPERVISED BY|DIRECTED BY)\s*[:.]?\s*\n*([A-Za-z0-9\s.,\-]+)', text, re.IGNORECASE)
    if sup_match:
        lines = [l.strip() for l in sup_match.group(1).split('\n') if l.strip()]
        if lines:
            meta.supervisors = [lines[0]]
            if len(lines) > 1 and any(rk in lines[1].upper() for rk in ["PROFESSOR", "LECTURER", "RANK", "DR"]):
                meta.supervisor_ranks = [lines[1].strip("() ")]

    # 5. Degree & Option
    if "MTECH" in text.upper() or "MASTER OF TECHNOLOGY" in text.upper():
        meta.degree = "Master of Technology"
        meta.degree_code = "MTech"
    elif "BTECH" in text.upper() or "BACHELOR OF TECHNOLOGY" in text.upper():
        meta.degree = "Bachelor of Technology"
        meta.degree_code = "BTech"
    elif "HND" in text.upper() or "HIGHER NATIONAL DIPLOMA" in text.upper():
        meta.degree = "Higher National Diploma"
        meta.degree_code = "HND"
    elif "PH.D" in text.upper() or "DOCTOR OF PHILOSOPHY" in text.upper():
        meta.degree = "Doctor of Philosophy"
        meta.degree_code = "PhD"
    elif "MSC" in text.upper() or "MASTER OF SCIENCE" in text.upper():
        meta.degree = "Master of Science"
        meta.degree_code = "MSc"
    elif "BSC" in text.upper() or "BACHELOR OF SCIENCE" in text.upper():
        meta.degree = "Bachelor of Science"
        meta.degree_code = "BSc"

    # Search for Department
    dept_match = re.search(r'Department of\s+([A-Za-z\s&]+?)(?:in the|of the|\n|,|\()', text, re.IGNORECASE)
    if dept_match:
        d = dept_match.group(1).strip()
        if len(d) > 3 and "UNIVERSITY" not in d.upper():
            meta.department = d
            meta.option = d

    # Search for Option / Specialization explicitly
    opt_match = re.search(r'(?:Option|Specialization|Speciality|Field|Branch)\s*[:.]?\s*([A-Za-z\s&]+?)(?:\n|,|\(|\.|$)', text, re.IGNORECASE)
    raw_opt = opt_match.group(1).strip() if opt_match else None

    # Search for School / Faculty
    for code, info in UBA_ESTABLISHMENTS.items():
        if info["name_en"] in text.upper() or f"({info['code']})" in text.upper():
            meta.faculty = info["name_en"]
            meta.faculty_code = info["code"]
            meta.motto = info["motto"]
            break

    # Resolve Department vs Option (prevents option from replacing department)
    dept, opt = resolve_department_and_option(meta.faculty_code, meta.department, raw_opt or meta.option)
    meta.department = dept
    meta.option = opt

    # Month & Year
    date_match = re.search(r'\b(JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\s+(20\d\d)\b', text, re.IGNORECASE)
    if date_match:
        meta.submission_month = date_match.group(1).upper()
        meta.submission_year = date_match.group(2)

    return meta


class ParsedDocument:
    def __init__(self):
        self.source_path: str = ""
        self.paragraphs: List[Dict[str, Any]] = []
        self.headings: List[Dict[str, Any]] = []
        self.tables_count: int = 0
        self.figures_count: int = 0
        self.page_count: int = 1
        self.raw_text: str = ""
        self.tables_text: List[str] = []
        self.metadata: DocumentMetadata = DocumentMetadata()
        self.margins: Dict[str, float] = {}  # in cm
        self.fonts_used: List[str] = []
        self.line_spacings: List[float] = []
        self.has_toc: bool = False
        self.has_declaration: bool = False
        self.has_certification: bool = False
        self.has_abstract: bool = False
        self.has_resume: bool = False
        self.has_acknowledgements: bool = False
        self.has_references: bool = False
        self.has_appendices: bool = False
        self.extracted_tables: List[Dict[str, Any]] = []


def parse_docx(file_path: str) -> ParsedDocument:
    """Parses a Word .docx document into structured representation."""
    doc = docx.Document(file_path)
    parsed = ParsedDocument()
    parsed.source_path = os.path.abspath(file_path)

    if doc.sections:
        s = doc.sections[0]
        parsed.margins = {
            "top": round(s.top_margin.cm, 2) if s.top_margin else 2.54,
            "bottom": round(s.bottom_margin.cm, 2) if s.bottom_margin else 2.54,
            "left": round(s.left_margin.cm, 2) if s.left_margin else 2.54,
            "right": round(s.right_margin.cm, 2) if s.right_margin else 2.54,
        }

    all_text = []
    fonts_set = set()
    spacings_set = set()

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue

        all_text.append(txt)
        style_name = p.style.name if p.style else "Normal"
        
        is_heading = False
        level = 0
        if "Heading 1" in style_name or "Chapter" in txt[:20] or (p.runs and p.runs[0].bold and p.runs[0].font.size and p.runs[0].font.size.pt >= 14):
            is_heading = True
            level = 1
        elif "Heading 2" in style_name or re.match(r'^\d+\.\d+\s', txt):
            is_heading = True
            level = 2
        elif "Heading 3" in style_name or re.match(r'^\d+\.\d+\.\d+\s', txt):
            is_heading = True
            level = 3

        for run in p.runs:
            if run.font.name:
                fonts_set.add(run.font.name)
        if p.paragraph_format.line_spacing:
            try:
                spacings_set.add(float(p.paragraph_format.line_spacing))
            except (ValueError, TypeError):
                pass

        para_info = {
            "text": txt,
            "style": style_name,
            "is_heading": is_heading,
            "level": level,
            "bold": any(r.bold for r in p.runs),
            "alignment": str(p.alignment) if p.alignment else "LEFT"
        }
        parsed.paragraphs.append(para_info)
        if is_heading:
            parsed.headings.append(para_info)

        txt_upper = txt.upper()
        if "TABLE OF CONTENTS" in txt_upper:
            parsed.has_toc = True
        if "DECLARATION" in txt_upper:
            parsed.has_declaration = True
        if "CERTIFICATION" in txt_upper:
            parsed.has_certification = True
        if "ABSTRACT" in txt_upper:
            parsed.has_abstract = True
        if "RÉSUMÉ" in txt_upper or "RESUME" in txt_upper:
            parsed.has_resume = True
        if "ACKNOWLEDGEMENT" in txt_upper or "ACKNOWLEDGMENTS" in txt_upper:
            parsed.has_acknowledgements = True
        if "REFERENCES" in txt_upper or "BIBLIOGRAPHY" in txt_upper:
            parsed.has_references = True
        if "APPENDIX" in txt_upper or "APPENDICES" in txt_upper or "ANNEX" in txt_upper:
            parsed.has_appendices = True

    parsed.tables_count = len(doc.tables)
    for t_idx, t in enumerate(doc.tables):
        t_data = []
        for row in t.rows:
            row_cells = []
            for cell in row.cells:
                c_text = cell.text.strip()
                row_cells.append(c_text)
                if c_text:
                    parsed.tables_text.append(c_text)
            t_data.append(row_cells)
        parsed.extracted_tables.append({
            "index": t_idx,
            "rows": t_data,
            "num_rows": len(t.rows),
            "num_cols": len(t.columns) if t.columns else 0
        })

    img_count = 0
    for part in doc.part.related_parts.values():
        if "image" in part.content_type:
            img_count += 1
    parsed.figures_count = img_count

    parsed.raw_text = "\n".join(all_text)
    parsed.fonts_used = list(fonts_set)
    parsed.line_spacings = list(spacings_set)
    words = len(parsed.raw_text.split())
    parsed.page_count = max(1, words // 350)
    parsed.metadata = extract_metadata_from_text_and_tables(parsed.raw_text, parsed.tables_text)

    return parsed


def parse_pdf(file_path: str) -> ParsedDocument:
    """Parses a PDF file into structured representation."""
    parsed = ParsedDocument()
    parsed.source_path = os.path.abspath(file_path)
    reader = PdfReader(file_path)
    parsed.page_count = len(reader.pages)

    all_text = []
    for i, page in enumerate(reader.pages):
        txt = page.extract_text() or ""
        lines = [l.strip() for l in txt.split("\n") if l.strip()]
        for line in lines:
            all_text.append(line)
            is_heading = False
            level = 0
            if re.match(r'^(CHAPTER\s+\d+|CHAPITRE\s+\d+)', line, re.IGNORECASE):
                is_heading = True
                level = 1
            elif re.match(r'^\d+\.\d+\s+[A-Z]', line):
                is_heading = True
                level = 2
            elif re.match(r'^\d+\.\d+\.\d+\s+[A-Z]', line):
                is_heading = True
                level = 3

            para_info = {
                "text": line,
                "is_heading": is_heading,
                "level": level,
                "page": i + 1
            }
            parsed.paragraphs.append(para_info)
            if is_heading:
                parsed.headings.append(para_info)

            line_upper = line.upper()
            if "TABLE OF CONTENTS" in line_upper:
                parsed.has_toc = True
            if "DECLARATION" in line_upper:
                parsed.has_declaration = True
            if "CERTIFICATION" in line_upper:
                parsed.has_certification = True
            if "ABSTRACT" in line_upper:
                parsed.has_abstract = True
            if "RÉSUMÉ" in line_upper or "RESUME" in line_upper:
                parsed.has_resume = True
            if "ACKNOWLEDGEMENT" in line_upper:
                parsed.has_acknowledgements = True
            if "REFERENCES" in line_upper:
                parsed.has_references = True

    parsed.raw_text = "\n".join(all_text)
    parsed.margins = {"left": 2.54, "right": 2.54, "top": 2.54, "bottom": 2.54}
    parsed.metadata = extract_metadata_from_text_and_tables(parsed.raw_text)
    return parsed


def parse_document(file_path: str) -> ParsedDocument:
    """Universal parser accepting .docx or .pdf."""
    path_lower = file_path.lower()
    if path_lower.endswith(".docx") or path_lower.endswith(".docx.bak"):
        return parse_docx(file_path)
    elif path_lower.endswith(".pdf"):
        return parse_pdf(file_path)
    else:
        raise ValueError(f"Unsupported file format: {os.path.splitext(file_path)[1]}. Please provide .docx or .pdf.")
